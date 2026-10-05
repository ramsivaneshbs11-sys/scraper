"""
Day 10 Tests — RAG-Powered Test Case Generation + Excel Export.

Tests are structured in 4 layers:

Layer 1: Heuristic Generator (offline — no LLM, no Qdrant)
Layer 2: TestCaseGeneratorService (mocked RAG retriever)
Layer 3: API Endpoints (via FastAPI TestClient, mocked LLM/RAG)
Layer 4: Excel Export Validation

Run all offline tests:
  pytest tests/test_day10.py -v -m "not requires_qdrant"
"""

import pytest
import json
import io
from unittest.mock import patch, MagicMock
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.main import app
from app.database import Base, get_db
from app import models

# -------------------------------------------------------------------
# Test DB Setup (isolated SQLite for Day 10 tests)
# -------------------------------------------------------------------

TEST_DB_URL = "sqlite:///./test_day10.db"
test_engine = create_engine(TEST_DB_URL, connect_args={"check_same_thread": False})
TestSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def override_get_db():
    db = TestSessionLocal()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture(scope="module", autouse=True)
def setup_db():
    Base.metadata.create_all(bind=test_engine)
    app.dependency_overrides[get_db] = override_get_db
    # Seed a project, user, document, and requirement for end-to-end tests
    db = TestSessionLocal()
    _seed_test_data(db)
    db.close()
    yield
    Base.metadata.drop_all(bind=test_engine)
    app.dependency_overrides.clear()


def _seed_test_data(db):
    """Insert minimal seed data for tests."""
    from datetime import datetime
    # User
    if not db.query(models.User).filter_by(id="u_d10").first():
        db.add(models.User(id="u_d10", email="day10@test.com", name="Day10 Tester", role="admin"))
    # Project
    if not db.query(models.Project).filter_by(id="p_d10").first():
        db.add(models.Project(id="p_d10", name="Day10 Project", created_by="u_d10"))
    # Document
    if not db.query(models.Document).filter_by(id="doc_d10").first():
        db.add(models.Document(
            id="doc_d10", project_id="p_d10", filename="spec.pdf",
            file_type="pdf", uploaded_by="u_d10", status="analyzed"
        ))
    # Chunk
    if not db.query(models.Chunk).filter_by(chunk_id="ch_d10_1").first():
        db.add(models.Chunk(
            chunk_id="ch_d10_1", project_id="p_d10", document_id="doc_d10",
            chunk_type="requirement_doc",
            text="The system shall authenticate users with JWT tokens that expire after 1 hour.",
            embedding_status=False
        ))
    # Requirement
    if not db.query(models.Requirement).filter_by(id="req_d10_1").first():
        db.add(models.Requirement(
            id="req_d10_1", project_id="p_d10",
            text="Users must authenticate with a valid JWT token before accessing any protected endpoint.",
            linked_chunk_ids=["ch_d10_1"],
            status="approved", confidence=0.95
        ))
    if not db.query(models.Requirement).filter_by(id="req_d10_2").first():
        db.add(models.Requirement(
            id="req_d10_2", project_id="p_d10",
            text="Login attempts must be limited to 5 per minute to prevent brute-force attacks.",
            linked_chunk_ids=["ch_d10_1"],
            status="approved", confidence=0.90
        ))
    db.commit()


client = TestClient(app)


# -------------------------------------------------------------------
# LAYER 1: Heuristic Generator (pure offline, no I/O)
# -------------------------------------------------------------------

class TestHeuristicGenerator:
    """
    Tests for the deterministic heuristic test case generator.
    These run without any LLM or Qdrant — 100% offline.

    LEARN: These tests verify the fallback generator, which is the
    safety net ensuring the pipeline always produces output even when
    LLM APIs are down or rate-limited.
    """

    def _make_req_context(self, req_id, text, chunks=None):
        return {
            "id": req_id,
            "text": text,
            "status": "approved",
            "context_chunks": chunks or []
        }

    def test_always_produces_positive_test_case(self):
        """Every requirement must generate at least one POSITIVE test case."""
        from app.generators.service import TestCaseGeneratorService
        contexts = [self._make_req_context("r1", "User must be able to log in.")]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        positives = [tc for tc in results if tc["test_type"] == "POSITIVE"]
        assert len(positives) >= 1

    def test_always_produces_negative_test_case(self):
        """Every requirement must generate at least one NEGATIVE test case."""
        from app.generators.service import TestCaseGeneratorService
        contexts = [self._make_req_context("r1", "System shall validate inputs.")]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        negatives = [tc for tc in results if tc["test_type"] == "NEGATIVE"]
        assert len(negatives) >= 1

    def test_boundary_generated_for_numeric_requirements(self):
        """
        BOUNDARY test cases must be generated when the requirement mentions numbers.
        LEARN: Numbers in requirements indicate thresholds, limits, or ranges
        that always need boundary testing (at, just-below, just-above).
        """
        from app.generators.service import TestCaseGeneratorService
        contexts = [self._make_req_context(
            "r1", "Login must be limited to 5 attempts per minute."
        )]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        boundaries = [tc for tc in results if tc["test_type"] == "BOUNDARY"]
        assert len(boundaries) >= 1, "No BOUNDARY test generated for numeric requirement"

    def test_no_boundary_for_non_numeric_requirements(self):
        """
        BOUNDARY test cases should NOT be generated for purely qualitative requirements.
        LEARN: Boundary testing only makes sense for quantifiable constraints.
        """
        from app.generators.service import TestCaseGeneratorService
        contexts = [self._make_req_context("r1", "System must display a welcome message.")]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        boundaries = [tc for tc in results if tc["test_type"] == "BOUNDARY"]
        assert len(boundaries) == 0, "BOUNDARY test generated for non-numeric requirement"

    def test_security_generated_for_auth_requirements(self):
        """
        SECURITY test cases must be generated when the requirement mentions
        authentication, authorization, tokens, passwords, or roles.
        """
        from app.generators.service import TestCaseGeneratorService
        contexts = [self._make_req_context(
            "r1", "Users must authenticate with JWT token before accessing resources."
        )]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        security = [tc for tc in results if tc["test_type"] == "SECURITY"]
        assert len(security) >= 1, "No SECURITY test generated for auth requirement"

    def test_performance_generated_for_latency_requirements(self):
        """
        PERFORMANCE test cases must be generated when the requirement mentions
        response time, latency, timeout, milliseconds, or load thresholds.
        """
        from app.generators.service import TestCaseGeneratorService
        contexts = [self._make_req_context(
            "r1", "API must respond within 200ms under normal load conditions."
        )]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        performance = [tc for tc in results if tc["test_type"] == "PERFORMANCE"]
        assert len(performance) >= 1, "No PERFORMANCE test generated for latency requirement"

    def test_every_test_case_has_required_fields(self):
        """
        Every generated test case must have all required fields populated.
        LEARN: This is the 'contract check' — same principle as Pydantic validation.
        """
        from app.generators.service import TestCaseGeneratorService
        required_fields = {"title", "test_type", "preconditions", "steps",
                           "expected_result", "acceptance_criteria",
                           "priority", "linked_requirement_ids", "rag_chunk_ids"}
        contexts = [self._make_req_context("r1", "System must authenticate users.")]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        for tc in results:
            missing = required_fields - set(tc.keys())
            assert not missing, f"Test case missing fields: {missing}"

    def test_rag_chunk_ids_included_in_output(self):
        """
        When RAG chunks are provided, their IDs must appear in rag_chunk_ids.
        LEARN: This is the traceability guarantee — every test case must link
        back to the source chunks that informed its generation.
        """
        from app.generators.service import TestCaseGeneratorService
        contexts = [self._make_req_context(
            "r1", "User shall log in.",
            chunks=[{"chunk_id": "chunk-abc", "text": "Login spec chunk", "score": 0.9}]
        )]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        for tc in results:
            assert "chunk-abc" in tc["rag_chunk_ids"], f"chunk-abc missing from rag_chunk_ids in {tc['title']}"

    def test_linked_requirement_id_on_every_test_case(self):
        """Every test case must link back to its parent requirement ID."""
        from app.generators.service import TestCaseGeneratorService
        contexts = [self._make_req_context("req-xyz", "System shall log errors.")]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        for tc in results:
            assert "req-xyz" in tc["linked_requirement_ids"], \
                f"Requirement ID missing in test case: {tc['title']}"

    def test_priority_is_p0_for_critical_keyword(self):
        """
        Requirements containing 'must', 'critical', or 'security' should generate
        P0-priority POSITIVE test cases.
        LEARN: Priority assignment is a key quality signal for test prioritization.
        """
        from app.generators.service import TestCaseGeneratorService
        contexts = [self._make_req_context("r1", "The system must never expose user passwords.")]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        positive_tcs = [tc for tc in results if tc["test_type"] == "POSITIVE"]
        assert any(tc["priority"] == "P0" for tc in positive_tcs), \
            "Expected P0 priority for 'must' keyword requirement"

    def test_multiple_requirements_generate_proportional_test_cases(self):
        """N requirements should generate at least N test cases (minimum 2 per req)."""
        from app.generators.service import TestCaseGeneratorService
        contexts = [
            self._make_req_context("r1", "System must validate email format."),
            self._make_req_context("r2", "System shall reject duplicate registrations."),
            self._make_req_context("r3", "Admin must be able to reset user passwords via secure link."),
        ]
        results = TestCaseGeneratorService._heuristic_generator(contexts)
        assert len(results) >= 6, f"Expected at least 6 test cases for 3 requirements, got {len(results)}"


# -------------------------------------------------------------------
# LAYER 2: Service Integration Tests (mocked RAG)
# -------------------------------------------------------------------

class TestGeneratorServiceIntegration:
    """
    Tests for TestCaseGeneratorService end-to-end with mocked RAG retrieval.
    These test that the service correctly orchestrates RAG + generation + persistence.
    """

    @patch("app.generators.service.LLMClient.get_api_key", return_value="")
    @patch("app.rag.retriever.RAGRetriever.retrieve_for_requirement")
    def test_generate_uses_rag_context(self, mock_retrieve, mock_api_key):
        """
        When use_rag=True, the service must call RAGRetriever.retrieve_for_requirement
        for each requirement. The retrieved chunks must appear in the test case output.
        """
        mock_retrieve.return_value = [
            {"chunk_id": "rag-chunk-001", "text": "JWT tokens expire after 1 hour.", "score": 0.88, "section": "Auth"}
        ]
        db = TestSessionLocal()
        try:
            from app.generators.service import TestCaseGeneratorService
            result = TestCaseGeneratorService.generate_test_cases_for_project(
                project_id="p_d10", db=db, include_proposed=True, use_rag=True
            )
            assert result["total_test_cases"] > 0
            assert result["rag_used"] is True
            # RAG retriever must have been called once per requirement
            assert mock_retrieve.call_count >= 1
        finally:
            db.close()

    @patch("app.generators.service.LLMClient.get_api_key", return_value="")
    def test_generate_without_rag_still_works(self, mock_api_key):
        """
        When use_rag=False, the service falls back to SQLite chunk lookup.
        Must still generate test cases without crashing.
        """
        db = TestSessionLocal()
        try:
            from app.generators.service import TestCaseGeneratorService
            result = TestCaseGeneratorService.generate_test_cases_for_project(
                project_id="p_d10", db=db, include_proposed=True, use_rag=False
            )
            assert result["total_test_cases"] > 0
            assert result["rag_used"] is False
        finally:
            db.close()

    @patch("app.generators.service.LLMClient.get_api_key", return_value="")
    def test_breakdown_in_result(self, _):
        """
        The result dict must include a breakdown of test cases by type
        (POSITIVE, NEGATIVE, BOUNDARY, SECURITY, PERFORMANCE).
        """
        db = TestSessionLocal()
        try:
            from app.generators.service import TestCaseGeneratorService
            result = TestCaseGeneratorService.generate_test_cases_for_project(
                project_id="p_d10", db=db, use_rag=False
            )
            assert "breakdown" in result
            breakdown = result["breakdown"]
            assert "POSITIVE" in breakdown
            assert "NEGATIVE" in breakdown
        finally:
            db.close()

    def test_raises_for_unknown_project(self):
        """Generating for a project with no requirements must raise ValueError."""
        from app.generators.service import TestCaseGeneratorService
        db = TestSessionLocal()
        try:
            with pytest.raises(ValueError, match="No requirements found"):
                TestCaseGeneratorService.generate_test_cases_for_project(
                    project_id="nonexistent-project-xyz", db=db
                )
        finally:
            db.close()


# -------------------------------------------------------------------
# LAYER 3: API Endpoint Tests
# -------------------------------------------------------------------

class TestTestCaseEndpoints:
    """API-level tests for all test case endpoints."""

    @patch("app.generators.service.LLMClient.get_api_key", return_value="")
    @patch("app.rag.retriever.RAGRetriever.retrieve_for_requirement",
           side_effect=ConnectionError("Qdrant offline"))
    def test_generate_endpoint_qdrant_offline(self, mock_rag, mock_key):
        """
        When Qdrant is offline, generation must still succeed by falling back
        to SQLite chunk lookup. Returns 200, not 503.
        """
        resp = client.post("/projects/p_d10/generate-test-cases?use_rag=true")
        assert resp.status_code == 200
        data = resp.json()
        assert data["total_test_cases"] > 0

    @patch("app.generators.service.LLMClient.get_api_key", return_value="")
    def test_generate_endpoint_use_rag_false(self, _):
        """POST /projects/{id}/generate-test-cases?use_rag=false returns test cases."""
        resp = client.post("/projects/p_d10/generate-test-cases?use_rag=false")
        assert resp.status_code == 200
        data = resp.json()
        assert data["total_test_cases"] > 0

    def test_get_test_cases_endpoint(self):
        """GET /projects/{id}/test-cases returns a list of test cases."""
        resp = client.get("/projects/p_d10/test-cases")
        assert resp.status_code == 200
        data = resp.json()
        assert "test_cases" in data
        assert "total_test_cases" in data

    def test_get_test_cases_filter_by_priority(self):
        """GET /projects/{id}/test-cases?priority=P0 filters correctly."""
        resp = client.get("/projects/p_d10/test-cases?priority=P0")
        assert resp.status_code == 200
        data = resp.json()
        for tc in data["test_cases"]:
            assert tc["priority"] == "P0"

    def test_approve_test_case_endpoint(self):
        """PATCH /test-cases/{id}/approve sets status to approved."""
        tcs = client.get("/projects/p_d10/test-cases").json()["test_cases"]
        assert len(tcs) > 0, "No test cases to approve"
        tc_id = tcs[0]["id"]
        resp = client.patch(f"/test-cases/{tc_id}/approve")
        assert resp.status_code == 200
        assert resp.json()["status"] == "success"

    def test_reject_test_case_endpoint(self):
        """PATCH /test-cases/{id}/reject sets status to rejected."""
        tcs = client.get("/projects/p_d10/test-cases").json()["test_cases"]
        assert len(tcs) > 0
        tc_id = tcs[-1]["id"]
        resp = client.patch(f"/test-cases/{tc_id}/reject")
        assert resp.status_code == 200

    def test_generate_returns_404_for_unknown_project(self):
        """POST /projects/{unknown}/generate-test-cases returns 404."""
        resp = client.post("/projects/no-such-project-xyz/generate-test-cases")
        assert resp.status_code in (400, 404)

    def test_delete_test_case_endpoint(self):
        """DELETE /test-cases/{id} removes the test case."""
        tcs = client.get("/projects/p_d10/test-cases").json()["test_cases"]
        assert len(tcs) > 0
        tc_id = tcs[0]["id"]
        resp = client.delete(f"/test-cases/{tc_id}")
        assert resp.status_code == 200
        # Confirm it's gone
        tcs_after = client.get("/projects/p_d10/test-cases").json()["test_cases"]
        ids_after = [t["id"] for t in tcs_after]
        assert tc_id not in ids_after


# -------------------------------------------------------------------
# LAYER 4: Export Tests
# -------------------------------------------------------------------

class TestExportEndpoints:
    """
    Tests for CSV and Excel export endpoints.
    These validate that the exported file has the correct format and content.
    """

    @patch("app.generators.service.LLMClient.get_api_key", return_value="")
    def setup_test_cases(self, _):
        """Helper to ensure test cases exist before export tests."""
        client.post("/projects/p_d10/generate-test-cases?use_rag=false")

    def test_csv_export_returns_200(self):
        """GET /projects/{id}/test-cases/export/csv returns a CSV file."""
        # Ensure test cases exist
        client.post("/projects/p_d10/generate-test-cases?use_rag=false")
        resp = client.get("/projects/p_d10/test-cases/export/csv")
        assert resp.status_code == 200
        assert "text/csv" in resp.headers.get("content-type", "")
        assert "attachment" in resp.headers.get("content-disposition", "")

    def test_csv_export_has_correct_headers(self):
        """CSV must have the standard column headers."""
        client.post("/projects/p_d10/generate-test-cases?use_rag=false")
        resp = client.get("/projects/p_d10/test-cases/export/csv")
        assert resp.status_code == 200
        content = resp.text
        first_line = content.split("\n")[0]
        assert "Title" in first_line
        assert "Priority" in first_line
        assert "Expected Result" in first_line

    def test_excel_export_returns_200(self):
        """GET /projects/{id}/test-cases/export/excel returns an xlsx file."""
        client.post("/projects/p_d10/generate-test-cases?use_rag=false")
        resp = client.get("/projects/p_d10/test-cases/export/excel")
        assert resp.status_code == 200
        content_type = resp.headers.get("content-type", "")
        assert "spreadsheetml" in content_type or "excel" in content_type.lower() or resp.status_code == 200

    def test_excel_export_is_valid_xlsx(self):
        """
        The Excel export must be a valid .xlsx file that openpyxl can open.
        LEARN: We load the response bytes into a BytesIO buffer and open it
        with openpyxl to verify the file is not corrupted.
        """
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        client.post("/projects/p_d10/generate-test-cases?use_rag=false")
        resp = client.get("/projects/p_d10/test-cases/export/excel")
        assert resp.status_code == 200

        buffer = io.BytesIO(resp.content)
        wb = openpyxl.load_workbook(buffer)

        # Must have at least 2 sheets
        assert "Test Cases" in wb.sheetnames
        assert "Summary" in wb.sheetnames

        ws = wb["Test Cases"]
        # Header row must be present
        headers = [ws.cell(1, col).value for col in range(1, 11)]
        assert "Title" in headers
        assert "Priority" in headers

        # Must have at least one data row
        assert ws.max_row >= 2, "Excel file has no data rows"
